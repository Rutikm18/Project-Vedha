# Builds margin_model.xlsx: price list + partner margin model. Example numbers only; founders replace the yellow cells.
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.comments import Comment
wb=Workbook(); ws=wb.active; ws.title="Margin model"
A="Arial"
f=lambda **k: Font(name=A, **k)
BLUE=f(color="0000FF"); BOLD=f(bold=True); NORM=f(); TITLE=f(bold=True,size=14); RED=f(bold=True,color="C00000")
YEL=PatternFill("solid",fgColor="FFFF00"); HEAD=PatternFill("solid",fgColor="DCE3EC")
INR='₹#,##0;(₹#,##0);-'; PCT='0.0%;(0.0%);-'; NUM='#,##0;(#,##0);-'
thin=Side(style="thin",color="C5D0DE")
def put(cell,v,font=NORM,fmt=None,fill=None):
    c=ws[cell]; c.value=v; c.font=font
    if fmt: c.number_format=fmt
    if fill: c.fill=fill
    return c
for col,w in zip("ABCDEF",[46,18,18,18,16,58]): ws.column_dimensions[col].width=w
put("A1","Vedha: 7-day check price list and partner margin model",TITLE)
put("A2","EXAMPLE NUMBERS, NOT A PROPOSAL. Replace every yellow cell with your own figures, then sign off.",RED)
put("A3","Legend: blue text on yellow = input you edit. Black = formula (don't edit). INR, excluding GST unless stated.",f(italic=True))
def header(r,labels):
    for i,l in enumerate(labels):
        c=ws.cell(row=r,column=i+1,value=l); c.font=BOLD; c.fill=HEAD
header(5,["PRICE LIST (per check)","Starter","Standard","Plus","Blended","Notes"])
put("A6","Office size covered")
for c,v in zip("BCD",["Up to 25 devices","Up to 75 devices","Up to 150 devices"]): put(f"{c}6",v,BLUE,fill=YEL)
put("F6","EXAMPLE tiers. [CONFIRM] device bands")
put("A7","List price per check (₹, excl. GST)")
for c,v in zip("BCD",[15000,30000,50000]): put(f"{c}7",v,BLUE,INR,YEL)
put("F7","EXAMPLE prices, not a proposal")
put("A8","Tier mix (share of clients)")
for c,v in zip("BCD",[0.5,0.35,0.15]): put(f"{c}8",v,BLUE,PCT,YEL)
put("E8","=SUM(B8:D8)",NORM,PCT); put("F8","Must total 100% (check the Blended cell)")
header(10,["ASSUMPTIONS","Value","","","","Notes / source"])
rows=[("GST rate",0.18,PCT,"18% on IT/consulting services (SAC 9983). Verify with your CA."),
("Partner discount on list price",0.30,PCT,"EXAMPLE. The partner's check margin"),
("Founder hours per check",6,NUM,"Replace with the measured hours from #16 (Partner Zero checks)"),
("Founder cost per hour (₹)",1500,INR,"EXAMPLE opportunity cost"),
("Infrastructure + tooling per check (₹)",500,INR,"EXAMPLE: hosting, probe, report generation"),
("Remediation hours per client (average)",8,NUM,"EXAMPLE: partner work to fix findings"),
("Partner hourly rate (₹)",1500,INR,"EXAMPLE: the partner's own rate"),
("Share of clients who buy remediation",0.6,PCT,"EXAMPLE"),
("Renewal rate (year 2 check)",0.5,PCT,"EXAMPLE"),
("Renewal price (% of list)",0.8,PCT,"EXAMPLE")]
r=11; refs=[]
for lab,v,fmt,note in rows:
    put(f"A{r}",lab); put(f"B{r}",v,BLUE,fmt,YEL); put(f"F{r}",note); refs.append(f"$B${r}"); r+=1
G,DISC,HRS,COST,INF,RH,RATE,ATT,REN,RENP=refs
header(22,["PER CLIENT","Starter","Standard","Plus","Blended","How it's calculated"])
lines=[
("Customer pays (incl. GST)", lambda c:f"={c}7*(1+{G})", INR, "List × (1 + GST)"),
("Partner buys from Vedha (excl. GST)", lambda c:f"={c}7*(1-{DISC})", INR, "List × (1 − partner discount)"),
("Partner margin on the check", lambda c:f"={c}7-{c}24", INR, "List − partner buy price"),
("Vedha delivery cost", lambda c:f"={HRS}*{COST}+{INF}", INR, "Founder hours × cost + infra"),
("Vedha net per check", lambda c:f"={c}24-{c}26", INR, "Partner buy price − delivery cost"),
("Vedha net margin", lambda c:f"=IF({c}24=0,0,{c}27/{c}24)", PCT, "Net ÷ Vedha revenue"),
("Partner remediation revenue (expected)", lambda c:f"={RH}*{RATE}*{ATT}", INR, "Hours × rate × share who buy"),
("Partner renewal margin (expected, year 2)", lambda c:f"={REN}*{c}7*{RENP}*{DISC}", INR, "Renewal rate × renewal price × discount"),
("Partner total per client", lambda c:f"={c}25+{c}29+{c}30", INR, "Check margin + remediation + renewal"),
]
for i,(lab,fn,fmt,note) in enumerate(lines):
    rr=23+i; put(f"A{rr}",lab, BOLD if lab.startswith(("Partner total","Vedha net per")) else NORM)
    for c in "BCD": put(f"{c}{rr}",fn(c),NORM,fmt)
    put(f"E{rr}", "=IF(E24=0,0,E27/E24)" if fmt==PCT else f"=SUMPRODUCT($B$8:$D$8,B{rr}:D{rr})", NORM, fmt)
    put(f"F{rr}",note)
header(33,["PER PARTNER PER YEAR","Scenario A","Scenario B","Scenario C","","Uses the Blended column"])
put("A34","Clients per partner per year")
for c,v in zip("BCD",[5,10,25]): put(f"{c}34",v,BLUE,NUM,YEL)
for rr,lab,b in [(35,"Partner earnings per year",True),(36,"Partner earnings per month",False),(37,"Vedha net from this partner per year",True),(38,"Founder hours this partner needs per year",False)]:
    put(f"A{rr}",lab,BOLD if b else NORM)
for c in "BCD":
    put(f"{c}35",f"={c}34*$E$31",NORM,INR); put(f"{c}36",f"={c}35/12",NORM,INR)
    put(f"{c}37",f"={c}34*$E$27",NORM,INR); put(f"{c}38",f"={c}34*{HRS}",NORM,NUM)
put("F38","If this exceeds founder capacity, the price or the hours must change")
put("A40","Sign-off",BOLD); put("A41","Approved by: ____________   Date: ____________")
put("A42","Then copy the final tiers into the deck (slide 8) and the competitive analysis (#8).")
for row in ws.iter_rows(min_row=5,max_row=38,max_col=6):
    for c in row: c.border=Border(bottom=thin)
ws["B7"].comment=Comment("EXAMPLE only. The founders set the real prices.","Claude")
ws.freeze_panes="A5"
ws.page_setup.orientation="landscape"; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1
# No cached values (written by openpyxl): make Excel / Sheets compute everything on open.
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save("margin_model.xlsx"); print("saved")
