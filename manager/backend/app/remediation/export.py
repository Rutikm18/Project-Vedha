"""
export.py — the remediation tracker the customer's IT team works from.

Three views of the SAME live data the UI and the report show:
  * By fix              — one row per fix bundle (FIX ref, target, due, progress)
  * By host checklist   — one row per (host, bundle)
  * Findings            — one row per finding with its bundle, due date and any
                          extension (reason + planned date → SEBI "reason for
                          non-closure" / "planned closure date")

XLSX: frozen header row, autofilter, real date cells (project timezone, IST by
default — Excel has no timezone, so the header says which). CSV: a zip holding
the three sheets as UTF-8 CSV files, or one sheet as a single CSV.

Titles, hosts and banners are scanner-derived — a hostile service can make a
finding title start with "=". Neither format may hand that to a spreadsheet as
a formula (CSV/formula injection): XLSX cells are forced to the string type,
CSV cells get the OWASP leading-apostrophe guard.
"""
from __future__ import annotations

import csv
import io
import zipfile
from datetime import datetime
from typing import Any

from app.services.project_time import to_project_tz

TZ_LABEL = "IST"
SHEETS = ("By fix", "By host checklist", "Findings")
# URL/query key → sheet name, for single-sheet CSV downloads.
SHEET_KEYS = {"by_fix": "By fix", "by_host": "By host checklist", "findings": "Findings"}
_FORMULA_LEAD = ("=", "+", "-", "@", "\t", "\r")


def _local(value: Any) -> datetime | None:
    if isinstance(value, str) and value:
        try:
            value = datetime.fromisoformat(value)
        except ValueError:
            return None
    if not isinstance(value, datetime):
        return None
    return to_project_tz(value).replace(tzinfo=None)


def _enum(v: Any) -> Any:
    return getattr(v, "value", v)


def build_rows(bundles: list[dict], items: list[tuple[Any, Any, dict]],
               extensions: list[Any]) -> dict[str, tuple[list[str], list[list[Any]]]]:
    """Header + rows per sheet. Values are str / int / naive local datetime / None."""
    by_fix_h = ["FIX ref", "Fix", "Remediation type", "Target state", "Urgency", "Status",
                f"Due ({TZ_LABEL})", "Due source", "Open findings", "Closed findings", "Hosts",
                "Effort", "Change risk", "Verification check", "Expected result"]
    by_fix = [[b["ref"], b["title"], b["remediation_type"], b["target_state"], b["urgency"], b["status"],
               _local(b["due_at"]), b["due_source"], b["open_count"], b["closed_count"], len(b["hosts"]),
               b["effort"], b["remediation_risk"], b["verification_check"]["check"],
               b["verification_check"]["expected"]] for b in bundles]

    host_h = ["Host", "IP", "OS", "FIX ref", "Fix", "Ports", "Finding IDs", "Status",
              f"Due ({TZ_LABEL})", "Due source", f"Closed ({TZ_LABEL})"]
    host_rows = []
    for b in bundles:
        for h in b["hosts"]:
            host_rows.append([h["host"], h["ip"], h["os"], b["ref"], b["title"],
                              ", ".join(str(p) for p in h["ports"]), ", ".join(h["finding_ids"]), h["status"],
                              _local(h["due_at"]), h["due_source"], _local(h.get("closed_at"))])
    host_rows.sort(key=lambda r: (r[1], r[3]))

    fix_of: dict[str, dict] = {}
    for b in bundles:
        for h in b["hosts"]:
            for fid in h["finding_ids"]:
                fix_of[fid] = b
    ext_of: dict[str, Any] = {}
    for req in sorted(extensions, key=lambda r: getattr(r, "requested_at", None) or datetime.min):
        if req.status not in ("pending", "approved"):
            continue
        for fid in req.member_finding_ids or []:
            ext_of[str(fid)] = req

    find_h = ["Finding ID", "Title", "Severity", "Status", "Host", "IP", "Port", "FIX ref",
              "Remediation type", "Target state", f"Due ({TZ_LABEL})", "Due source",
              "Extension status", "Extension reason", f"Planned closure ({TZ_LABEL})", "Verification check"]
    find_rows = []
    for f, ctx, due in items:
        fid = str(f.id)
        b = fix_of.get(fid)
        ext = ext_of.get(fid)
        planned = (ext.decided_due_at or ext.requested_due_at) if ext is not None else None
        find_rows.append([
            fid, f.title, str(_enum(f.severity)), str(_enum(f.status)), ctx.text("host"), ctx.text("ip"),
            ctx.port.value if ctx.port.recorded else None, b["ref"] if b else None,
            b["remediation_type"] if b else None, b["target_state"] if b else None,
            _local((due or {}).get("at")), (due or {}).get("source"),
            ext.status if ext is not None else None, ext.reason if ext is not None else None,
            _local(planned), b["verification_check"]["check"] if b else None,
        ])
    sev_rank = {"critical": 0, "high": 1, "medium": 2, "low": 3, "info": 4}
    find_rows.sort(key=lambda r: (sev_rank.get(r[2], 5), r[1]))
    return {"By fix": (by_fix_h, by_fix), "By host checklist": (host_h, host_rows), "Findings": (find_h, find_rows)}


def to_xlsx(sheets: dict[str, tuple[list[str], list[list[Any]]]], *, title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    for name in SHEETS:
        header, rows = sheets[name]
        ws = wb.create_sheet(name)
        ws.append(header)
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in rows:
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_LEAD):
                    cell.data_type = "s"     # openpyxl would otherwise write it as a formula
        for col_idx, head in enumerate(header, start=1):
            letter = get_column_letter(col_idx)
            if "(" + TZ_LABEL + ")" in head:
                for cell in ws[letter][1:]:
                    cell.number_format = "yyyy-mm-dd hh:mm"
            width = max([len(str(head))] + [len(str(r[col_idx - 1] or "")) for r in rows[:500]])
            ws.column_dimensions[letter].width = min(max(12, width + 2), 60)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
    wb.properties.title = title
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv_cell(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, str) and v.startswith(_FORMULA_LEAD):
        return "'" + v
    return v


def _csv_text(header: list[str], rows: list[list[Any]]) -> str:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(header)
    for r in rows:
        w.writerow([_csv_cell(v) for v in r])
    return "\ufeff" + out.getvalue()      # BOM so Excel reads UTF-8


def to_csv(sheets: dict[str, tuple[list[str], list[list[Any]]]], sheet: str) -> bytes:
    """One sheet as a single UTF-8 CSV file."""
    header, rows = sheets[sheet]
    return _csv_text(header, rows).encode("utf-8")


def to_csv_zip(sheets: dict[str, tuple[list[str], list[list[Any]]]]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name in SHEETS:
            header, rows = sheets[name]
            zf.writestr(name.lower().replace(" ", "_") + ".csv", _csv_text(header, rows))
    return buf.getvalue()
